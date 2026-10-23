#!/usr/bin/env python3
"""Export original King Rail V2 biological source as Word and source-SVG figure proofs.

No science refitting or reassignment of nonmigrants. No author list, journal
endorsement, ethics permissions or proof of movement causality is supplied.
"""
from __future__ import annotations
from pathlib import Path
import hashlib,json,subprocess,shutil,zipfile
import pypandoc,cairosvg,fitz
from docx import Document
from docx.shared import Pt,Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
V1=ROOT/"manuscript/LOUIS_HYDROLOGICAL_BUFFERING_MANUSCRIPT_V1.md"
V2=ROOT/"manuscript/LOUIS_KING_RAIL_MAIN_V2_AUTHOR_REVIEW.md"
CONTRACT=ROOT/"manuscript/MANUSCRIPT_NUMERIC_CONTRACT_V2.json"
FIG_SOURCE=ROOT/"manuscript/FIGURE_DATA_CONTRACT_V2.json"
FIGDIR=ROOT/"manuscript/rendered_figures"
OUT=ROOT/"results/generated_king_rail_v2_review"
ZIP=ROOT/"results/LOUIS_KING_RAIL_V2_AUTHOR_REVIEW_PACKAGE.zip"

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def convert(md,out,resources):
    src=out.with_suffix(".input.md")
    src.parent.mkdir(parents=True,exist_ok=True)
    src.write_text(md,encoding="utf-8")
    pypandoc.convert_file(str(src),"docx",
       format="markdown-yaml_metadata_block+tex_math_dollars",outputfile=str(out),
       extra_args=["--standalone","--resource-path="+":".join(str(p) for p in resources)])
    src.unlink(missing_ok=True)
    d=Document(str(out))
    for sty in ("Normal","Body Text","Heading 1","Heading 2","Heading 3"):
        if sty not in d.styles:continue
        s=d.styles[sty]
        s.font.name="Times New Roman"
        s.font.size=Pt({"Heading 1":12,"Heading 2":11,"Heading 3":10}.get(sty,10))
        if sty in ("Normal","Body Text"):
            s.paragraph_format.space_after=Pt(5)
            s.paragraph_format.line_spacing=1.42
        else:
            s.font.bold=True
            s.paragraph_format.keep_with_next=True
            s.paragraph_format.space_before=Pt(11)
    for section in d.sections:
        section.left_margin=Inches(.9)
        section.right_margin=Inches(.88)
        section.top_margin=Inches(.84)
        section.bottom_margin=Inches(.8)
        section.header_distance=Inches(.4)
        section.footer_distance=Inches(.41)
        p=section.header.paragraphs[0]
        p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        p.text="KING RAIL V2  |  AUTHOR REVIEW — NOT SUBMITTED"
        for run in p.runs:run.font.size=Pt(7)
        page=section.footer.paragraphs[0]
        page.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        page.add_run("Page ").font.size=Pt(8)
        f=OxmlElement("w:fldSimple")
        f.set(qn("w:instr"),"PAGE")
        page._p.append(f)
        for e in section._sectPr.findall(qn("w:lnNumType")):section._sectPr.remove(e)
        ln=OxmlElement("w:lnNumType")
        for k,v in (("countBy","1"),("start","1"),("distance","360"),("restart","continuous")):
            ln.set(qn("w:"+k),v)
        section._sectPr.append(ln)
    d.core_properties.author="Contributor metadata not yet supplied"
    d.core_properties.subject="Original King Rail reanalysis - author-review export"
    d.save(str(out))
    return out

def pdf_from(doc,out):
    render=OUT/"render"
    render.mkdir(exist_ok=True)
    z=subprocess.run(["libreoffice","-env:UserInstallation=file:///tmp/lo_louis_kingrail_review",
                      "--headless","--convert-to","pdf","--outdir",str(render),str(doc)],
                     capture_output=True,text=True,timeout=180)
    made=render/(doc.stem+".pdf")
    if z.returncode or not made.is_file():
        raise RuntimeError("DOCX render failed: "+z.stdout[:300]+" "+z.stderr[:300])
    shutil.copy2(made,out)
    return out

def proof(pdf,terms,minimum_pages,n_images=0):
    f=fitz.open(str(pdf))
    assert f.page_count>=minimum_pages
    words="\n".join(p.get_text(sort=True) for p in f)
    for k in terms:
        assert k.lower() in words.lower(),(pdf.name,k)
    images=0
    for p in f:
        images+=len(p.get_images(full=True))
        for b in p.get_text("blocks"):
            if len(b)<5 or not isinstance(b[4],str):continue
            x0,y0,x1,y1=b[:4]
            assert x0>=-2 and y0>=-2 and x1<=p.rect.width+2 and y1<=p.rect.height+2,(pdf.name,p.number,b)
    assert images>=n_images,(pdf.name,images)
    return {"file":pdf.name,"pages":f.page_count,"source_figures_or_images":images,"text_clipping_blocks":0}

def main():
    v2=V2.read_text(encoding="utf-8")
    assert "**Authors:** [to be completed]" in v2
    assert "190 matched events" in v2 and "10 birds" in v2
    assert "0.886" in v2 and "83.8%" in v2 and "0.658" in v2
    assert "cannot be paired with coordinate-level movement trajectories" in v2
    assert json.loads(CONTRACT.read_text())["schema"]=="louis.manuscript_numeric_contract.v2"
    OUT.mkdir(parents=True,exist_ok=True)
    header="**AUTHOR REVIEW. Source study ethics and authorship must be verified. Not submitted.**\n\n"
    main=convert(header+v2,OUT/"KingRail_V2_Main_AuthorReview.docx",[V2.parent])
    figures=OUT/"figures"
    figures.mkdir(exist_ok=True)
    markdown=(
      "# King Rail source-verified main figures — author review\n\n"
      "**Figure captions and canonical effects are in the unchanged King Rail V1/V2 manuscript.**\n\n"
      "This book is a visual author proof, not independent science or approved article artwork.\n\n"
    )
    for n in range(1,6):
        fig=FIGDIR/f"Figure{n}.svg"
        assert fig.is_file()
        shutil.copy2(fig,figures/fig.name)
        png=figures/f"Figure{n}_proof.png"
        cairosvg.svg2png(url=str(fig),write_to=str(png),scale=1.5)
        assert png.stat().st_size>1200
        markdown+=f"## Figure {n}\n\n![Source Figure {n}](Figure{n}_proof.png){{ width=6.1in }}\n\n"
    figs=convert(markdown,OUT/"KingRail_V2_Figures_AuthorReview.docx",[figures])
    mp=pdf_from(main,OUT/"KingRail_V2_Main_PROOF.pdf")
    fp=pdf_from(figs,OUT/"KingRail_V2_FigureBook_AuthorReview.pdf")
    audit={
       "status":"PASS_AUTHOR_REVIEW_DOCUMENT_EXPORT_NOT_SUBMITTED",
       "source_sha256":{str(x.relative_to(ROOT)):sha(x) for x in (V1,V2,CONTRACT,FIG_SOURCE)},
       "main":proof(mp,["King Rails","190 matched events","0.886","0.658","83.8%","References"],8),
       "figure_book":proof(fp,["Figure 1","Figure 5","author review"],3,5),
       "line_number_sections":sum(
           s._sectPr.find(qn("w:lnNumType")) is not None for s in Document(str(main)).sections),
       "existing_figures":5,
       "human_review_remaining":[
           "actual authors and affiliations","original source animal-handling approvals",
           "journal-specific manuscript and abstract limits","complete reference proof and contributor approval",
           "author figure inspection and release DOI"],
    }
    assert audit["line_number_sections"]>=1
    (OUT/"SOURCE_AND_LAYOUT_QC.json").write_text(json.dumps(audit,indent=2)+"\n",encoding="utf-8")
    (OUT/"AUTHOR_REVIEW_CHECKLIST.md").write_text(
        "# King Rail V2 author-review checklist\n\n**Not submitted.**\n\n"+
        "\n".join("- [ ] "+x for x in audit["human_review_remaining"])+"\n",encoding="utf-8")
    shutil.rmtree(OUT/"render",ignore_errors=True)
    ZIP.parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(ZIP,"w",zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for f in sorted(OUT.rglob("*")):
            if f.is_file():z.write(f,f.relative_to(OUT))
    print(json.dumps({"status":audit["status"],"main":audit["main"],"figures":audit["figure_book"],
       "archive":str(ZIP),"zip_bytes":ZIP.stat().st_size,"sha256":sha(ZIP)},indent=2))

if __name__=="__main__":
    main()
